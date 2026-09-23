"""Presentation-only curves and structure diagram from frozen research outputs.

No fitting, new cohorts, threshold changes or scientific artifact replacement.
Run data with scientific dependencies, then slides with python-pptx 1.0.2.
"""
import argparse
import json
from pathlib import Path
from hashlib import sha256


ARMS={
 'DWM_binomial':('DWM + binomial shrinkage','#64748B'),
 'LOG21_binomial':('LOG21 + binomial shrinkage','#087F8C'),
 'LOG21_beta_binomial':('LOG21 + beta-binomial shrinkage','#9156A6'),
 'LOG21_count_tail':('LOG21 + existing count-tail score','#D47B22'),
}


def digest(path):
    with path.open('rb') as f:return sha256(f.read()).hexdigest()


def data(source,out):
    import numpy as np
    import pandas as pd
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from sklearn.metrics import roc_curve,precision_recall_curve,roc_auc_score,average_precision_score
    from verify import read_table
    out.mkdir(exist_ok=False,parents=True)
    frame=read_table(source/'tables/per_site_scores.tsv.gz')
    frame=frame[frame.cohort.eq('nested')&frame.depth.eq('full')&frame.method.isin(ARMS)]
    assert np.isfinite(frame.score).all()
    fold_source=read_table(source/'tables/fold_metrics.tsv')
    fold_source=fold_source[fold_source.cohort.eq('nested')&fold_source.depth.eq('full')]
    grid=np.linspace(0,1,1001);curves=[];metrics=[];max_error=0.
    for (cell,tf,library,fold),group in frame.groupby(['cell','tf','library','fold']):
        reference=None
        for method in ARMS:
            g=group[group.method.eq(method)].sort_values('site_id')
            ids=g[['site_id','chip_label']].reset_index(drop=True)
            if reference is None:reference=ids
            else:pd.testing.assert_frame_equal(reference,ids)
            if g.chip_label.nunique()!=2:continue
            y=g.chip_label.to_numpy(int);score=g.score.to_numpy(float)
            auc=roc_auc_score(y,score);ap=average_precision_score(y,score)
            saved=fold_source[fold_source.cell.eq(cell)&fold_source.tf.eq(tf)&fold_source.library.eq(library)&fold_source.fold.eq(fold)&fold_source.method.eq(method)]
            assert len(saved)==1
            error=max(abs(auc-saved.auroc.iloc[0]),abs(ap-saved.auprc.iloc[0]));max_error=max(max_error,error)
            assert error<1e-12
            fpr,tpr,_=roc_curve(y,score,drop_intermediate=False)
            precision,recall,_=precision_recall_curve(y,score)
            # Average precision integrates a step function, not a line between
            # recall jumps. Use the precision at the next attained recall.
            ascending=recall[::-1];p_ascending=precision[::-1]
            pr=p_ascending[np.searchsorted(ascending,grid,side='left')]
            pr[0]=1.;pr[-1]=y.mean()
            for name,values in [('ROC',np.interp(grid,fpr,tpr)),('PR',pr)]:
                curves.extend(dict(cell=cell,tf=tf,library=library,fold=fold,method=method,curve=name,x=float(x),y=float(v)) for x,v in zip(grid,values))
            metrics.append(dict(cell=cell,tf=tf,library=library,fold=fold,method=method,sites=len(g),prevalence=y.mean(),auroc=auc,average_precision=ap))
    curves=pd.DataFrame(curves);metrics=pd.DataFrame(metrics)
    metrics.to_csv(out/'fold_curve_metrics.tsv',sep='\t',index=False)
    curves.to_csv(out/'fold_curve_coordinates.tsv.gz',sep='\t',index=False)
    # Preserve the original equal-chromosome -> equal-library -> equal-task estimand.
    library=curves.groupby(['cell','tf','library','method','curve','x'],as_index=False).y.mean()
    task=library.groupby(['cell','tf','method','curve','x'],as_index=False).y.mean()
    mean=task.groupby(['method','curve','x'],as_index=False).y.mean()
    task.to_csv(out/'task_curve_coordinates.tsv.gz',sep='\t',index=False)
    mean.to_csv(out/'macro_curve_coordinates.tsv',sep='\t',index=False)
    task_metrics=metrics.groupby(['cell','tf','library','method'],as_index=False)[['auroc','average_precision','prevalence']].mean().groupby(['cell','tf','method'],as_index=False)[['auroc','average_precision','prevalence']].mean()
    summary=task_metrics.groupby('method')[['auroc','average_precision','prevalence']].mean()
    assert (task_metrics.groupby('method').size()==20).all()
    original=read_table(source/'tables/task_metrics.tsv')
    original=original[original.cohort.eq('nested')&original.depth.eq('full')&original.method.isin(ARMS)]
    expected=original.groupby(['cell','tf','method'],as_index=False)[['auroc','auprc']].mean().groupby('method')[['auroc','auprc']].mean()
    np.testing.assert_allclose(summary.auroc,expected.auroc,atol=1e-12,rtol=0)
    np.testing.assert_allclose(summary.average_precision,expected.auprc,atol=1e-12,rtol=0)
    summary.to_csv(out/'macro_metrics.tsv',sep='\t')
    for method in ARMS:
        for name,metric in [('ROC','auroc'),('PR','average_precision')]:
            g=mean[mean.method.eq(method)&mean.curve.eq(name)]
            assert abs(np.trapezoid(g.y,g.x)-summary.loc[method,metric])<.002
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False})
    for name in ['ROC','PR']:
        fig,ax=plt.subplots(figsize=(6.25,3.6))
        for method,(label,color) in ARMS.items():
            g=mean[mean.method.eq(method)&mean.curve.eq(name)]
            ax.plot(g.x,g.y,color=color,lw=2.3,label=label)
        if name=='ROC':
            ax.plot([0,1],[0,1],ls='--',color='#B8C1CA',lw=1,label='Chance reference')
            ax.set(xlabel='False-positive rate',ylabel='True-positive rate')
        else:
            ax.axhline(float(summary.prevalence.iloc[0]),ls='--',color='#B8C1CA',lw=1,label='Matched-cohort prevalence')
            ax.set(xlabel='Recall',ylabel='Precision')
        ax.set(xlim=(0,1),ylim=(0,1));ax.grid(alpha=.12);fig.tight_layout()
        fig.savefig(out/(name+'_macro.png'),dpi=220);fig.savefig(out/(name+'_macro.svg'));plt.close(fig)
    # A diagram of prescribed regions, not an observed biological profile.
    fig,ax=plt.subplots(figsize=(6.1,1.15))
    ax.plot([-45,45],[0,0],color='#64748B',lw=2)
    for a,b,color,label in [(-40,-12,'#93BBC3','Left shoulder'),(-7,7,'#D47B22','Center'),(12,40,'#93BBC3','Right shoulder')]:
        ax.fill_between([a,b],-.2,.2,color=color);ax.text((a+b)/2,.29,label,ha='center',fontsize=10)
    ax.set_xticks([-40,-12,-7,0,7,12,40]);ax.tick_params(axis='x',labelsize=9)
    ax.set(xlim=(-45,45),ylim=(-.35,.65),yticks=[],xlabel='Base pairs from motif center · geometry schematic')
    for spine in ax.spines.values():spine.set_visible(False)
    fig.tight_layout(pad=.3);fig.savefig(out/'geometry.png',dpi=220);plt.close(fig)
    provenance=dict(source_score_sha256=digest(source/'tables/per_site_scores.tsv.gz'),source_model_sha256=digest(source/'models/frozen_LOG21.npz'),
        source_fold_metrics_sha256=digest(source/'tables/fold_metrics.tsv'),available_tasks=20,original_tasks=21,unavailable='HepG2 HMGA1',
        initial_libraries=6,selected_site_method_rows=len(frame),fold_method_rows=len(metrics),maximum_metric_error=max_error,
        weighting='equal nonempty chromosome within library, equal library within task, equal available task',
        curves='ROC linear interpolation; PR recall-step function consistent with average precision; 1001 grid coordinates. Exact metric labels agree with the approximate curve areas within 0.002; no pooled scores.',
        common_support='identical site IDs and labels asserted for all four arms within every fold',
        inference='descriptive presentation of existing frozen predictions; no new fits, cohort choices, thresholds or qualification',
        meaning='matched-cohort precision is not genome-wide precision; no pooled cross-task score thresholds')
    (out/'CURVE_VERIFICATION.json').write_text(json.dumps(provenance,indent=2)+'\n')
    print(json.dumps(provenance,indent=2),flush=True)


def slides(out):
    from datetime import datetime,timezone
    import csv
    from pptx import Presentation
    from pptx.util import Inches
    from presentation import text,box,NAVY,GRAY,TEAL,RUST,PALE,WHITE
    with (out/'macro_metrics.tsv').open() as f:metrics={r['method']:r for r in csv.DictReader(f,delimiter='\t')}
    deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
    props=deck.core_properties;props.title='fp-tools: parametric model and measured development performance'
    props.author=props.last_modified_by='fp-tools research';props.created=props.modified=datetime.now(timezone.utc)
    slide=deck.slides.add_slide(deck.slide_layouts[6]);box(slide,0,0,13.34,.1,TEAL)
    text(slide,.4,.3,12.5,.5,'How the parametric model works',29,bold=True)
    text(slide,.4,.94,12.5,.43,'A frozen sequence-bias model supplies the expected cleavage pattern; a separate detector measures depletion.',16)
    cards=[(.4,'1  Oriented sequence','21 bases around each candidate cut.\nReverse-complement reverse strands.'),
           (4.74,'2  Log-linear bias score','Single-base effects + pairwise interactions\nat separations of 1, 5 and 9 bases.'),
           (9.08,'3  Expected cut pattern','Softmax converts scores to probabilities.\nMultiply by the observed regional total.')]
    for x,title,body in cards:
        box(slide,x,1.58,3.85,1.22,PALE);text(slide,x+.14,1.74,3.57,.3,title,18,bold=True)
        text(slide,x+.14,2.15,3.55,.57,body,13)
    text(slide,4.32,2.02,.35,.4,'→',24,color=TEAL,bold=True);text(slide,8.66,2.02,.35,.4,'→',24,color=TEAL,bold=True)
    text(slide,.55,2.99,12.2,.45,'ηⱼ = Σₖ aₖ(bⱼₖ) + Σd∈{1,5,9} Σₖ Bₖ,d(bⱼₖ, bⱼ,ₖ₊d)       πⱼ = exp(ηⱼ) / Σₗ exp(ηₗ)',19,bold=True)
    text(slide,.55,3.5,12.18,.69,'ηⱼ: log bias at position j; a: single-base coefficients; B: base-pair interaction coefficients; b: sequence base.\nFit conditional multinomial cut locations given total counts; squared-coefficient penalty = 0.0001. Frozen five-seed mitochondrial ensemble.',13,color=GRAY)
    box(slide,.4,4.4,6.13,2.38,PALE);box(slide,6.77,4.4,6.15,2.38,PALE)
    text(slide,.56,4.57,5.82,.3,'From sequence expectation to regional counts',17,bold=True)
    slide.shapes.add_picture(str(out/'geometry.png'),Inches(.6),Inches(4.98),width=Inches(5.7),height=Inches(1.06))
    text(slide,.6,6.15,5.7,.5,'For each strand: N = center + both shoulders; C = center cuts.\np = null center probability; expected center count = N × p.',12.5)
    text(slide,6.94,4.57,5.8,.3,'Separate signed-depletion model',17,bold=True)
    text(slide,6.96,5.05,5.75,.55,'logit(qₛ) = logit(pₛ) − θ;   θ ~ Normal(0, 1)\nCₛ | Nₛ, θ, ρ ~ Beta-binomial(Nₛ, qₛ, ρ)',16,bold=True)
    text(slide,6.96,5.82,5.7,.75,'Rank by posterior mean θ; positive means depletion.\nFallback: ρ = 0 (binomial). No counts → score 0, no call.\nFit ρ on naked DNA; exclude the evaluated chromosome.',12.5)
    text(slide,.45,7.01,12.3,.31,'LOG21 = 21-base log-linear model. Strand likelihood is a working model; posterior values are not binding probabilities.',12,color=RUST,bold=True)
    slide.notes_slide.notes_text_frame.text=(
        'Structure checked against src/fp_tools/tools/parametric_bias.py (BiasFeatureSpec.loglinear21 and log_scores), '
        'benchmarks/scripts/next_round/propensities.py and shrinkage.py. Frozen coefficient SHA-256 is in CURVE_VERIFICATION.json. '
        'The 21-position base terms use sum-to-zero constraints; pair matrices are double-centered for identifiability. '
        'The five-seed mitochondrial LOG21 ensemble is frozen, shifts +4/-5, L2=0.0001; no occupancy-label optimization or new bias training was done. '
        'Softmax normalizes over the region under consideration. For detector central probability, normalize the frozen sequence propensities over center and shoulders: center -7..7 inclusive, shoulders -40..-12 and 12..40 inclusive; gaps are excluded. '
        'N is conditioned on, not predicted by the bias model. C and N refer to each strand separately, with one shared signed site effect theta. '
        'The beta-binomial variance is N*q*(1-q)*[1+(N-1)*rho]. Rho is control-fitted within the permitted fitting chromosomes. '
        'The two strand likelihood factors are composite/working factors because fragment endpoints can be dependent. '
        'The geometry image is a schematic, not observed protection. Source references and the original evidence archive remain unchanged.')
    slide=deck.slides.add_slide(deck.slide_layouts[6]);box(slide,0,0,13.34,.1,TEAL)
    text(slide,.4,.28,12.5,.44,'Observed ranking performance: all available tasks',27,bold=True)
    text(slide,.4,.87,12.48,.26,'Full depth · K562 + HepG2 · 3 biological libraries each · held training chromosomes · 20/21 tasks (HMGA1 unavailable)',13.5)
    text(slide,.4,1.19,12.48,.22,'Dinucleotide weight matrix (DWM): same-detector propensity baseline, not a native TOBIAS benchmark.',11.5,color=GRAY)
    text(slide,.53,1.46,5.95,.3,'Receiver operating characteristic (ROC)',18,bold=True)
    text(slide,6.94,1.46,5.96,.3,'Precision–recall (PR)',18,bold=True)
    slide.shapes.add_picture(str(out/'ROC_macro.png'),Inches(.34),Inches(1.85),width=Inches(6.3),height=Inches(3.63))
    slide.shapes.add_picture(str(out/'PR_macro.png'),Inches(6.69),Inches(1.85),width=Inches(6.3),height=Inches(3.63))
    text(slide,.5,5.49,12.3,.28,'Area under the ROC curve (AUROC); area under the PR curve (AUPRC), computed as average precision.',11,color=GRAY)
    for i,(method,(label,color)) in enumerate(ARMS.items()):
        x=.48+(i%2)*6.43;y=5.91+(i//2)*.36
        box(slide,x,y+.07,.24,.11,color.lstrip('#'))
        r=metrics[method]
        text(slide,x+.34,y,5.77,.31,f"{label}   {float(r['auroc']):.3f} / {float(r['average_precision']):.3f}",12,bold=True)
    text(slide,.49,6.68,12.2,.27,'Same binomial detector: +0.0227 AUROC; +2.70% task-relative AUPRC. Macro average: chromosome → library → task.',11,color=TEAL,bold=True)
    box(slide,.4,7.06,12.5,.3,NAVY)
    text(slide,.52,7.085,12.24,.22,'Development only: coverage confounding and unmet calibration/shape gates prevent promotion. Metric pairs above: AUROC / AUPRC.',11.5,color=WHITE,bold=True)
    slide.notes_slide.notes_text_frame.text=(
        'Data source: complete per_site_scores.tsv.gz from the verified next-round evidence packet; nested cohort, full depth, all six libraries. '
        'All four arms use identical site IDs and labels within each nonempty chromosome fold, asserted before plotting. '
        'Receiver operating characteristic curves are linearly interpolated within each fold on 1001 grid coordinates. Precision–recall uses the recall-step function consistent with average precision, with endpoint convention precision=1 at recall=0 and prevalence at recall=1. '
        'Curves are averaged equally across nonempty chromosomes within each library, then across three libraries within each task, then across all 20 available tasks. '
        'The 21st task, HepG2 HMGA1, has no matched evaluation cohort and is not set to zero. '
        'Metric numbers are exact fold AUROC and average precision under the same weighting; they are not numerical areas of interpolated mean curves. '
        'PR dashed reference is the matched-cohort positive prevalence (0.5), not population prevalence. Matched-cohort precision is not genome-wide precision. '
        'The underlying LOG21 versus DWM binomial comparison is +0.022675 mean AUROC and +2.70356% mean task-relative average precision. '
        'A ratio of the displayed grand means is not the mean of task-relative changes. '
        'LOG21 binomial and beta-binomial lose -0.037118 and -0.043105 mean AUROC against the existing LOG21 count-tail scorer. '
        'Actual coverage imbalance remains in all 60 available task/library panels. The old count-tail score may exploit coverage confounding; its higher curve does not establish superior biological footprint detection. '
        'The same-site observed/expected profiles do not change when only the detector changes. The native full-depth DWM residual comparator remains unqualified. '
        'Neither shrinkage method met advancement rules, and no protected controls or new-context labels were opened. '
        'See CURVE_VERIFICATION.json and complete fold/task curve tables for reproducibility; no new scientific fitting was done for this presentation.')
    target=out/'fp_tools_model_and_curves.pptx';deck.save(target)
    reloaded=Presentation(target);assert len(reloaded.slides)==2
    for slide in reloaded.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=deck.slide_width+Inches(.02)
            assert shape.top+shape.height<=deck.slide_height
    print(target,flush=True)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path);p.add_argument('--outdir',type=Path,required=True);p.add_argument('--stage',choices=['data','slides'],required=True)
    a=p.parse_args()
    if a.stage=='data':data(a.source.resolve(),a.outdir.resolve())
    else:slides(a.outdir.resolve())
