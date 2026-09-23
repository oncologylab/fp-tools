"""Create a two-slide, editable development-results briefing from final tables."""
import argparse
import csv
from datetime import datetime, timezone
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


NAVY='142C43';TEAL='087F8C';RUST='A34431';GRAY='526476';PALE='EEF4F7';WHITE='FFFFFF'


def read(path):
    with path.open() as f:return list(csv.DictReader(f,delimiter='\t'))


def box(slide,x,y,w,h,fill):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill);s.line.fill.background()
    return s


def text(slide,x,y,w,h,value,size=15,color=NAVY,bold=False):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    f=s.text_frame;f.word_wrap=True
    f.margin_left=f.margin_right=0;f.margin_top=f.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=line
        p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(5)
    return s


def task_table(slide,rows,x,y,w,cell):
    text(slide,x,y,w,.3,cell+' · all recorded factors',16,bold=True)
    table=slide.shapes.add_table(len(rows)+1,3,Inches(x),Inches(y+.36),Inches(w),Inches(.24*(len(rows)+1))).table
    for col,width in zip(table.columns,[w*.42,w*.26,w*.32]):col.width=Inches(width)
    values=[['Factor','Δ AUROC','Relative AUPRC']]
    for r in rows:
        valid=r['status']=='DEVELOPMENT_DIAGNOSTIC'
        values.append([r['tf'],f"{float(r['auroc_gain']):+.4f}" if valid else 'Unavailable',
                       f"{float(r['relative_auprc_gain']):+.2%}" if valid else 'No cohort'])
    for i,row in enumerate(values):
        for j,value in enumerate(row):
            c=table.cell(i,j);c.text=value;c.margin_top=c.margin_bottom=Inches(.01)
            c.margin_left=Inches(.08);c.margin_right=Inches(.04)
            c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string(NAVY if i==0 else PALE if i%2 else WHITE)
            p=c.text_frame.paragraphs[0];p.font.name='Arial';p.font.size=Pt(11)
            p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string(WHITE if i==0 else NAVY)


def execute(out):
    summary=read(out/'tables/full_screen_summary.tsv');tasks=read(out/'tables/full_screen_task_deltas.tsv')
    def select(method,reference):return next(r for r in summary if r['method']==method and r['reference']==reference)
    gain=select('LOG21_binomial','DWM_binomial')
    deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
    deck.core_properties.title='fp-tools: measured development gains and qualification limits'
    deck.core_properties.author='fp-tools research'
    deck.core_properties.last_modified_by='fp-tools research'
    deck.core_properties.created=datetime.now(timezone.utc)
    deck.core_properties.modified=datetime.now(timezone.utc)
    deck.core_properties.subject='Bounded September 2026 development investigation; not a release or promotion'
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    box(slide,0,0,13.334,.12,TEAL)
    text(slide,.45,.32,12.4,.45,'Frozen parametric bias: modest ranking gains, uneven across factors',25,bold=True)
    text(slide,.45,.92,12.4,.5,'Same binomial-shrinkage detector and held sites: 21-base log-linear bias (LOG21) versus dinucleotide weight matrix (DWM) propensities.',14)
    text(slide,.45,1.46,12.4,.25,'Full depth · 6 biological libraries · training chromosomes 1–15 · 20 available / 21 original tasks',14,color=GRAY)
    text(slide,.45,1.78,12.4,.18,'Area under the receiver operating characteristic curve (AUROC); area under the precision–recall curve (AUPRC), measured as average precision.',10.2,color=GRAY)
    for x,value,label in [(.45,f"{float(gain['mean_auroc_gain']):+.4f}",'Mean AUROC change'),(4.64,f"{float(gain['mean_relative_auprc_gain']):+.2%}",'Mean relative AUPRC change'),(8.82,'No promotion','Development comparison only')]:
        box(slide,x,1.98,4.03,.84,PALE);text(slide,x+.15,2.06,3.7,.35,value,25,color=TEAL,bold=True)
        text(slide,x+.15,2.48,3.7,.22,label,12)
    selected=[r for r in tasks if r['method']=='LOG21_binomial' and r['reference']=='DWM_binomial']
    task_table(slide,[r for r in selected if r['cell']=='K562'],.45,3.02,5.97,'K562')
    task_table(slide,[r for r in selected if r['cell']=='HepG2'],6.91,3.02,5.97,'HepG2')
    text(slide,.48,6.51,12.35,.29,'Equal chromosome → equal biological library → equal available task. Missing tasks are not zeros. CTCF denotes CCCTC-binding factor.',10.2,color=GRAY)
    text(slide,.48,6.89,12.3,.4,'Ranking gain ≠ stronger footprint protection. This is an fp-tools propensity comparison, not an end-to-end TOBIAS benchmark.',12,color=RUST,bold=True)
    slide.notes_slide.notes_text_frame.text=(
        'Sources: tables/full_screen_summary.tsv and full_screen_task_deltas.tsv. All task rows are shown, including unavailable HepG2 HMGA1. '
        'Equal chromosome, then equal biological library, then equal available task weighting. Missing tasks are not zeros. '
        'LOG21 is the frozen five-seed mitochondrial ensemble, shifts +4/-5, L2=0.0001. The binomial signed-depletion model has a fixed Normal(0,1) prior. '
        'The conditional fixed-prediction AUROC difference interval is '+gain['fixed_prediction_bootstrap_low']+' to '+gain['fixed_prediction_bootstrap_high']+
        ', based on 1958 defined of 2000 synchronized chromosome resamples. It is not the qualified principal promotion interval or full-refit uncertainty. '
        'Actual full-depth count imbalance exists in every available task/library panel. Covariate baseline discrimination is substantial. '
        'CCCTC-binding factor (CTCF) is a positive control. Gene symbols are retained as source task identifiers. '
        'Mean +0.0227 and relative +2.70% do not meet the original +0.03 and +10% thresholds; moreover, this comparator is not the full native DWM residual arm required for promotion.')
    slide=deck.slides.add_slide(deck.slide_layouts[6]);box(slide,0,0,13.334,.12,RUST)
    text(slide,.45,.32,12.4,.48,'The new shrinkage scorers did not justify advancement',27,bold=True)
    text(slide,.45,.98,12.35,.5,'Primary: beta-binomial depletion with zero-centered shrinkage. Fallback: the same model with dispersion fixed at zero.',15)
    text(slide,.45,1.76,6.25,.4,'Compared with the existing LOG21 count-tail scorer',17,bold=True)
    for y,method,name in [(2.35,'LOG21_binomial','Binomial fallback'),(3.33,'LOG21_beta_binomial','Beta-binomial primary')]:
        r=select(method,'LOG21_count_tail');box(slide,.45,y,6.1,.83,PALE)
        text(slide,.62,y+.08,2.7,.25,name,16,bold=True)
        text(slide,3.44,y+.06,2.95,.3,f"AUROC {float(r['mean_auroc_gain']):+.4f}",20,color=RUST,bold=True)
        text(slide,3.44,y+.45,2.95,.25,f"Relative AUPRC {float(r['mean_relative_auprc_gain']):+.2%}",12,color=RUST)
    extra=select('LOG21_beta_binomial','LOG21_binomial')
    text(slide,.48,4.46,6.05,.7,f"Extra dispersion did not help: AUROC {float(extra['mean_auroc_gain']):+.4f} versus the simpler fallback.",17)
    text(slide,.48,5.4,6.08,.91,'These ranking losses do not prove an intrinsically worse biological detector: coverage confounding remains.',17,color=RUST,bold=True)
    box(slide,6.88,1.73,5.99,4.68,PALE)
    text(slide,7.07,1.94,5.62,.35,'What the evidence supports',19,bold=True)
    text(slide,7.07,2.52,5.56,3.8,
        '60/60 task–library panels retain coverage imbalance; count-related discrimination is not isolated.\n'
        '7/21 control panels show a central-mean mismatch; 18/21 checks fail for each candidate.*\n'
        'Changing only the detector leaves the same-site aggregate profiles unchanged.\n'
        'Final controls, full native comparator, functional uncertainty and resource gates remain unqualified.',16)
    text(slide,.48,6.59,12.34,.25,'* Examined naked-DNA replicate 1; exploratory intervals and dependent-site checks, not untouched safety evidence.',10.5,color=GRAY)
    box(slide,.45,7.00,12.42,.34,NAVY)
    text(slide,.6,7.035,12.1,.25,'Decision: stop this candidate search. Keep production unchanged and protected labels closed.',14,color=WHITE,bold=True)
    slide.notes_slide.notes_text_frame.text=(
        'Sources: full_screen_summary.tsv, actual_library_covariate_balance.tsv, covariate_baseline_task_metrics.tsv, '
        'control_mean_variance.tsv.gz, paired_control_rates.tsv, shape_metrics.tsv and promotion_gates.tsv. '
        'All twelve initial 10M/full endpoints completed across three libraries in each cell context, seed 2026. '
        'The full-depth screen was the only advancement screen; expensive additional depths/seeds and formal resources were intentionally not run. '
        'Coverage imbalance means absolute standardized mean difference >0.1 in log1p whole-window counts. '
        'Central-mean intervals are exploratory chromosome-block intervals, without multiplicity adjustment. '
        'Each candidate has 18 failing task-level informative-rate/Wilson-upper-bound/same-operator paired-DWM development checks, out of 21. '
        'Only examined naked-DNA replicate 1 was used. These checks do not substitute for original native-DWM or final untouched safety gates. '
        'No new stronger measured protection is inferred from ranking changes. Low-count non-detection remains partly informational and biologically unresolved. '
        'No new SK-N-SH/GM23338 labels or final naked-DNA replicate 3 were opened. See NEXT_DEEP_RESEARCH_PROMPT.md for falsifiable follow-up review questions.')
    target=out/'figures/fp_tools_development_summary.pptx';deck.save(target)
    restored=Presentation(target)
    assert len(restored.slides)==2
    for s in restored.slides:
        for shape in s.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=restored.slide_width+Inches(.01)
            assert shape.top+shape.height<=restored.slide_height
    print(target,flush=True)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--outdir',type=Path,required=True)
    execute(p.parse_args().outdir.resolve())
