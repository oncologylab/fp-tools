"""Question-led technical slides; no scientific scoring or fitting."""
import argparse
from pathlib import Path
from datetime import datetime,timezone
import csv
from pptx import Presentation
from pptx.util import Inches
from presentation import text,box,NAVY,TEAL,RUST,GRAY,PALE,WHITE


def execute(curves,out):
    out.mkdir(parents=True,exist_ok=True)
    deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
    deck.core_properties.title='fp-tools parametric bias: training, resolution and reuse'
    deck.core_properties.author=deck.core_properties.last_modified_by='fp-tools research'
    deck.core_properties.created=deck.core_properties.modified=datetime.now(timezone.utc)
    s=deck.slides.add_slide(deck.slide_layouts[6]);box(s,0,0,13.34,.1,TEAL)
    text(s,.4,.28,12.5,.47,'21-base input context → one prediction at every base',28,bold=True)
    text(s,.42,.89,12.4,.32,'21-base log-linear model (LOG21): shared sequence-bias coefficients, not a separate binding model for each factor.',15)
    for x,w,title in [(.4,3.6,'TRAIN'),(4.25,4.34,'MODEL'),(8.84,4.09,'APPLY')]:
        box(s,x,1.52,w,4.83,PALE);text(s,x+.17,1.7,w-.34,.3,title,20,color=TEAL,bold=True)
    text(s,.6,2.2,3.18,.7,'K562 + HepG2\nMitochondrial cuts + hg38',17,bold=True)
    text(s,.6,3.14,3.18,.85,'200-base windows\n65 train + 16 validation per cell\n22.78 million training cut events',15)
    text(s,.6,4.21,3.18,.8,'Target: cut counts at each base\nNo transcription-factor binding labels',15)
    text(s,.6,5.32,3.18,.72,'Fit five seeded models\nAverage coefficients; freeze',16,bold=True)
    # Sliding input contexts: 21 bases per output, shifted by one base.
    seq='ACGTGACCTAGCTACGTAGCTAC';x0=4.46;dx=.172
    for row,start in [(0,0),(1,1)]:
        y=2.32+row*.58
        for i in range(21):
            fill=TEAL if i==10 else 'D4E2EA'
            box(s,x0+i*dx,y,dx-.01,.26,fill)
            text(s,x0+i*dx+.013,y+.025,dx,.19,seq[i+start],9,WHITE if i==10 else NAVY,bold=i==10)
        text(s,4.49,y+.3,3.8,.2,'predict at j' if row==0 else 'shift 1 base → predict at j + 1',11,color=GRAY)
    text(s,4.47,3.68,3.91,.74,'21 single-base terms\n+ interactions at offsets 1, 5, 9',17,bold=True)
    text(s,4.47,4.6,3.9,.76,'η = sum of learned coefficients\nπⱼ = exp(ηⱼ) / Σ exp(η)',17,bold=True)
    text(s,4.47,5.6,3.9,.48,'Conditional multinomial fit\nSquared-weight penalty: 0.0001',13.5)
    text(s,9.03,2.2,3.68,.64,'New genomic sequence\n→ relative cut probabilities π',18,bold=True)
    text(s,9.03,3.18,3.68,.73,'Observed local cut total N\n→ expected count at base j: N × πⱼ',16)
    text(s,9.03,4.16,3.68,.24,'Illustration: three positions',13,color=GRAY)
    for i,(p,c,h) in enumerate([('.25','5',.35),('.50','10',.7),('.25','5',.35)]):
        x=9.42+i*1.01
        box(s,x,5.25-h,.62,h,TEAL)
        text(s,x,5.36,.75,.25,p,14,bold=True)
        text(s,x,5.75,.75,.25,c,14,bold=True)
    text(s,9.03,5.36,.38,.24,'π',14);text(s,9.03,5.75,.38,.24,'E',14)
    text(s,9.03,6.06,3.65,.19,'N = 20 → expected counts [5, 10, 5]',11)
    text(s,.47,6.62,12.24,.29,'Three different scales: 21 bases = sequence context • 1 base = prediction spacing • 15 bases = detector center (−7…+7)',14,bold=True)
    text(s,.47,7.08,12.24,.23,'The bias profile is not a footprint or binding probability. Single-base spacing does not establish single-base biological accuracy.',12,color=RUST)
    s.notes_slide.notes_text_frame.text=(
        'Verified historical training provenance: ENCODE K562 ENCFF834HWH and HepG2 ENCFF095VYO unfiltered alignment files retaining mitochondrial reads, hg38 reference. '
        'The active frozen coefficient file SHA-256 is 620f232f634423b8753e4d77f8cf72df87482c09193cb5e7e6b038917c47e89a. '
        'The 200-base cache has 81 mitochondrial windows per cell: 65 training and 16 validation, partitioned by a coordinate hash (mitochondria cannot use nuclear chromosome splits). '
        'K562 has 14516077 training cuts; HepG2 has 8267931; sum 22784008. These are historical cut-event counts, not independently counted complete fragments. '
        'The fit records 260 strand-window rows: (65+65)*2 strands, not 260 distinct genomic loci. Cells share mitochondrial coordinates and are not independent sequences. '
        'Read filters exclude unmapped, secondary, supplementary, quality-failed and duplicate reads, with minimum mapping quality 1. Cut convention is +4/-5 and reverse contexts are reverse-complemented. '
        'Five member runs 2026–2030 use the same full-depth training source, 100 epochs, learning rate 0.03, L2=0.0001. Their stored coefficients are averaged equally; five complete training runs are not repeated for each new transcription factor. '
        'The input has 21 nucleotide positions, 84 stored main effects, and 48 pair-position combinations at offsets 1,5,9 with 16 stored coefficients each; total 852 stored coefficients, subject to centering constraints. '
        'Log score equals the sum of single-base and pair-feature weights. Conditional multinomial likelihood fits the allocation of cuts given total coverage; it does not predict absolute sequencing depth. '
        'At application, calculate scores on a one-base grid; normalize over the chosen window/strand; multiply by its observed cut total. The bar illustration is arithmetic only, not experimental data. '
        'The scoring geometry has center -7..7 and shoulders -40..-12,12..40. The separate aggregate array spans 201 positions (-100..100), still with one-base spacing. '
        'This explanation checks existing records and code; it does not revalidate historical extraction, retrain the bias model or reopen protected inputs.')
    s=deck.slides.add_slide(deck.slide_layouts[6]);box(s,0,0,13.34,.1,TEAL)
    text(s,.4,.28,12.45,.48,'What is reused—and what changes for a new analysis?',27,bold=True)
    for x,title,body in [(.4,'Different transcription factor?','Same bias coefficients.\nOnly the motif locations change.'),
                         (4.66,'Different sample / project?','Same coefficients in this study.\nNew cut counts and site effects.'),
                         (8.92,'Different assay / protocol?','Transfer is not qualified.\nControl validation is unresolved.')]:
        box(s,x,1.04,4.0,1.17,PALE);text(s,x+.14,1.18,3.72,.3,title,16,bold=True)
        text(s,x+.14,1.61,3.72,.53,body,14)
    text(s,.48,2.36,12.2,.37,'The downstream detector is separate: estimate depletion θ per site; control-fit dispersion ρ; no per-factor prior.',14,bold=True)
    text(s,.48,2.88,5.9,.28,'Receiver operating characteristic (ROC)',15,bold=True)
    text(s,7.0,2.88,5.85,.28,'Precision–recall (PR)',15,bold=True)
    s.shapes.add_picture(str(curves/'ROC_macro.png'),Inches(.4),Inches(3.2),width=Inches(5.75),height=Inches(3.312))
    s.shapes.add_picture(str(curves/'PR_macro.png'),Inches(6.95),Inches(3.2),width=Inches(5.75),height=Inches(3.312))
    legend=[('64748B','Dinucleotide weight matrix + binomial'),('087F8C','LOG21 + binomial'),('9156A6','LOG21 + beta-binomial'),('D47B22','LOG21 + existing count-tail')]
    for i,(color,label) in enumerate(legend):
        x=.48+(i%2)*6.51;y=6.5+(i//2)*.25
        box(s,x,y+.045,.18,.09,color);text(s,x+.28,y,5.95,.21,label,11)
    text(s,.48,7.08,12.2,.23,'20/21 tasks, 6 libraries, full depth. Bias gains are modest; transfer/calibration remain unqualified. No method is promoted.',11.8,color=RUST,bold=True)
    s.notes_slide.notes_text_frame.text=(
        'LOG21 means 21-base log-linear sequence model. Bias coefficients do not accept transcription factor identity as an input; motif identities locate evaluation windows, and do not select different trained weights in this round. '
        'All six K562/HepG2 biological libraries reuse the same frozen mitochondrial bias coefficients. New observed counts change expected counts after depth conditioning and change inferred signed-depletion site effects. '
        'The beta-binomial downstream detector uses logit(q_s)=logit(p_s)-theta, theta~Normal(0,1); rank by posterior mean theta. Binomial fallback sets rho=0. '
        'Rho and calling thresholds were estimated on examined naked-DNA replicate-1 training chromosomes, leaving out the evaluated chromosome. They were not automatically learned from each new chromatin library, and their transfer is not established. '
        'A new project name is not a biological reason to retrain; nor do these results justify a promise of universal reuse across new sample types or assay protocols. No validated rule for retraining on arbitrary new projects has been established. '
        'Curves are unchanged from the independently checked presentation supplement: 20 available tasks, 3 biological libraries in each of K562 and HepG2, full depth, nested held-training-chromosome evaluation. HMGA1 in HepG2 is unavailable. '
        'Equal nonempty chromosome, then equal library, then equal task weighting. Four methods use identical site IDs per fold. All 2520 fold/method metrics reproduce exactly. PR uses the recall-step convention. Dashed lines show the chance ROC diagonal and matched-cohort precision prevalence 0.5. '
        'Mean area under the ROC curve / average precision: DWM-binomial 0.497423/0.580817; LOG21-binomial 0.520098/0.595556; LOG21-beta-binomial 0.514111/0.589434; LOG21 count-tail 0.557216/0.623919. '
        'The +0.022675 AUROC and +2.70356% task-relative average precision of LOG21 versus DWM under the same detector do not establish stronger canonical protection or native end-to-end TOBIAS superiority. '
        'All 60 available task/library panels retain actual coverage imbalance; null mean/calibration problems and original promotion gates remain unresolved. No sample-specific bias fit is authorized or performed here.')
    target=out/'fp_tools_training_prediction_reuse.pptx';deck.save(target)
    assert len(Presentation(target).slides)==2
    print(target)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--curves',type=Path,required=True);p.add_argument('--outdir',type=Path,required=True)
    a=p.parse_args();execute(a.curves.resolve(),a.outdir.resolve())
