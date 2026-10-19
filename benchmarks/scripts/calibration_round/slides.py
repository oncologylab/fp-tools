"""Two editable slides (method; defensible performance) from completed tables. Needs python-pptx."""
import argparse
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

INK, MUTED, ACCENT, BAD, LINE = (RGBColor(0x0b, 0x0b, 0x0b), RGBColor(0x52, 0x51, 0x4e), RGBColor(0x2a, 0x78, 0xd6),
                                 RGBColor(0xb0, 0x2a, 0x2a), RGBColor(0xc8, 0xc7, 0xc2))


def text(slide, x, y, w, h, lines, size=12, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame; frame.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        content, style = (line, {}) if isinstance(line, str) else line
        run = p.add_run(); run.text = content
        run.font.name = 'Arial'
        run.font.size = Pt(style.get('size', size)); run.font.bold = style.get('bold', bold)
        run.font.color.rgb = style.get('color', color); p.space_after = Pt(3)
    return box


def stage(slide, x, y, w, title, body, accent=False):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(1.35))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(0xe8, 0xf0, 0xfb) if accent else RGBColor(0xf3, 0xf2, 0xef)
    shape.line.color.rgb = ACCENT if accent else LINE
    frame = shape.text_frame; frame.word_wrap = True
    p = frame.paragraphs[0]; r = p.add_run(); r.text = title; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = INK
    p2 = frame.add_paragraph(); r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(9); r2.font.color.rgb = MUTED
    r.font.name = r2.font.name = 'Arial'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--outdir', type=Path, required=True)
    out = parser.parse_args().outdir.resolve()
    screen = pd.read_csv(out / 'tables/advancement_screen.tsv', sep='\t')
    prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    text(s, .5, .3, 12.3, .6, 'Method: can a label-free null correction make LOG21 calls trustworthy?', 22, bold=True)
    text(s, .5, .9, 12.3, .4, 'Development chr1–15 only · dinucleotide weight matrix (DWM) production default unchanged', 11, MUTED)
    steps = [('1 Observed cuts', 'Strand cuts from complete fragments in centre −7..7 and shoulders ±12..40 (n, c)', False),
             ('2 LOG21 (frozen)', '21-bp context → one propensity per base; trained on mitochondrial reads, no labels', False),
             ('3 Null recalibration', 'q = expit(a + b·logit p); a, b from naked-DNA rep 1, held chromosome out (candidate)', True),
             ('4 Count-tail test', 'P(central ≤ c | n₊, n₋, q); unchanged cap, rounding and 1.301 threshold', False),
             ('5 Call or abstain', 'CALL / NO_CALL; ABSTAIN if zero central cuts could not reach the threshold', False)]
    for i, (title, body, acc) in enumerate(steps):
        stage(s, .5 + i * 2.5, 1.5, 2.3, title, body, acc)
    text(s, .5, 3.05, 6.0, 3.6, [
        ('Why this candidate', {'bold': True, 'size': 13}),
        'LOG21 mispredicted naked-DNA central cleavage in 7/21 tasks, and the previous round modelled variance before fixing the mean.',
        ('Compared on identical sites and counts', {'bold': True, 'size': 13}),
        'Candidate (2 scalars) · intercept-only ablation · unchanged LOG21 reference · DWM propensities with the same detector · covariates only.',
        '6,020 held sites, 6 libraries × {10M, full}; 21 tasks (HepG2 HMGA1 unavailable).',
        ('Never used', {'bold': True, 'size': 13}),
        'Occupancy labels for bias fitting · SK-N-SH / GM23338 labels · naked-DNA replicate 3.'], 11)
    s.shapes.add_picture(str(out / 'figures/control_null_mean.png'), Inches(6.9), Inches(2.95), height=Inches(4.4))

    s = prs.slides.add_slide(blank)
    text(s, .5, .3, 12.3, .6, 'Calibration candidate rejected: ranking gains did not rescue protection', 22, bold=True)
    text(s, .5, .9, 12.3, .4, 'Single frozen screen, full depth, 20/21 tasks; DWM arm is a factorial control, not native DWM or TOBIAS',
         11, MUTED)
    s.shapes.add_picture(str(out / 'figures/roc_pr_average.png'), Inches(.4), Inches(1.4), width=Inches(6.6))
    lines = [('Supported', {'bold': True, 'size': 13, 'color': ACCENT}),
             'One intercept fixes the pooled naked-DNA mean (−16% under-prediction) and drops dispersion 1.22 → 1.04.',
             'CTCF protection reproduced: positive÷negative central O/E 0.69 / 0.67.',
             ('Not supported', {'bold': True, 'size': 13, 'color': BAD}),
             'Calibration: still 6/21 motif-specific discrepancies (gate ≤3).',
             'Accuracy: +0.037 AUROC [0.020, 0.053], +4.2% AP vs DWM (gate 10%); within coverage strata 0.513 = covariates 0.514.',
             'No canonical protection in any difficult family; 10M direction 4/6.',
             'ChIP-negative central O/E 0.34–0.99; calls 11.8% (neg) vs 18.3% (pos). Negative labels are not proven unbound controls.',
             ('Decision', {'bold': True, 'size': 13}),
             'Reject; keep LOG21 count-tail as research reference. Next only if authorized: naked-DNA-trained sequence model, '
             'tested on held naked-DNA chromosomes before any chromatin scoring.']
    text(s, 7.2, 1.35, 5.7, 5.6, lines, 10.5)
    gates = ' · '.join(f"{g.gate.split('_')[0]} {g.status.lower()}" for g in screen.itertuples() if g.gate[:2] in
                       {'A_', 'B_', 'C_', 'D_', 'E_', 'F_', 'G_', 'H_'})
    text(s, .5, 6.75, 12.3, .4, 'Gates: ' + gates + ' · CPU/memory not run (screen failed) · untouched controls locked', 9, MUTED)
    prs.save(out / 'figures/calibration_round_two_slides.pptx')


if __name__ == '__main__':
    main()
